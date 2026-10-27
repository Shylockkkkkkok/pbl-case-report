# -*- coding: utf-8 -*-
"""
add_image_slide.py —— 给已有 PPT 插入一页「标题 + 大图 + 图注」

用法：
    python add_image_slide.py --deck X.pptx --image fig.png --title "典型体征与影像" \
        [--after "锚点文字"] [--caption "数据来源：…"] [--margin-top 1.15]

--after 指定插在该文字所在页之后；省略则追加到末尾。
会自动沿用版式配色（藏青标题 + 朱红竖条 + 灰色页脚），并重排全篇页码。
"""
import argparse
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

NAVY = RGBColor(0x17, 0x3D, 0x67)
RED = RGBColor(0xAF, 0x3D, 0x3D)
GREY = RGBColor(0x5A, 0x6B, 0x7B)
FONT = "Noto Sans CJK SC"

ap = argparse.ArgumentParser()
ap.add_argument("--deck", required=True)
ap.add_argument("--image", required=True)
ap.add_argument("--title", required=True)
ap.add_argument("--caption", default="")
ap.add_argument("--after", default="")
ap.add_argument("--footer", default="")
ap.add_argument("--margin-top", type=float, default=1.15)
ap.add_argument("--margin-bottom", type=float, default=6.55)
a = ap.parse_args()

prs = Presentation(a.deck)
layout = next((l for l in prs.slide_masters[0].slide_layouts
               if not any(sh.is_placeholder for sh in l.shapes)), prs.slide_masters[0].slide_layouts[-1])
W = prs.slide_width / 914400.0
H = prs.slide_height / 914400.0

s = prs.slides.add_slide(layout)
for shp in list(s.shapes):
    shp._element.getparent().remove(shp._element)

bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.25), Inches(0.20), Inches(0.075), Inches(0.42))
bar.fill.solid(); bar.fill.fore_color.rgb = RED
bar.line.fill.background(); bar.shadow.inherit = False

tb = s.shapes.add_textbox(Inches(0.42), Inches(0.14), Inches(12.4), Inches(0.55))
tb.text_frame.word_wrap = True
r = tb.text_frame.paragraphs[0].add_run()
r.text = a.title
r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = NAVY; r.font.name = FONT

im = Image.open(a.image)
ar = im.width / im.height
box_w, box_h = W - 1.1, a.margin_bottom - a.margin_top
dw, dh = box_w, box_w / ar
if dh > box_h:
    dh = box_h; dw = dh * ar
s.shapes.add_picture(a.image, Inches((W - dw) / 2), Inches(a.margin_top), Inches(dw), Inches(dh))

if a.caption:
    cb = s.shapes.add_textbox(Inches(0.55), Inches(a.margin_top + dh + 0.10), Inches(W - 1.1), Inches(0.32))
    p = cb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    cr = p.add_run(); cr.text = a.caption
    cr.font.size = Pt(11); cr.font.color.rgb = GREY; cr.font.name = FONT

if a.footer:
    fb = s.shapes.add_textbox(Inches(0.25), Inches(H - 0.26), Inches(8), Inches(0.22))
    fr = fb.text_frame.paragraphs[0].add_run(); fr.text = a.footer
    fr.font.size = Pt(10.5); fr.font.color.rgb = GREY; fr.font.name = FONT
    nb = s.shapes.add_textbox(Inches(12.55), Inches(H - 0.29), Inches(0.55), Inches(0.22))
    pn = nb.text_frame.paragraphs[0]; pn.alignment = PP_ALIGN.RIGHT
    nr = pn.add_run(); nr.text = "00"
    nr.font.size = Pt(10.5); nr.font.color.rgb = GREY; nr.font.name = FONT

if a.after:
    idx = len(prs.slides) - 1
    target = -1
    for i, sl in enumerate(prs.slides):
        if i == idx:
            continue
        if any(sh.has_text_frame and a.after in sh.text_frame.text for sh in sl.shapes):
            target = i
            break
    if target >= 0:
        lst = prs.slides._sldIdLst
        items = list(lst)
        el = items[idx]
        lst.remove(el)
        lst.insert(target + 1, el)

for i, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t.isdigit() and len(t) <= 2 and sh.left / 914400 > 11.5 and sh.top / 914400 > 6.9:
            for p in sh.text_frame.paragraphs:
                for r2 in p.runs:
                    r2.text = "%02d" % i

prs.save(a.deck)
print("added; slides:", len(prs.slides))
