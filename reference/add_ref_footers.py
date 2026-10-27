# -*- coding: utf-8 -*-
"""
add_ref_footers.py —— 给每页加页脚参考文献

用法：
    python add_ref_footers.py --deck X.pptx --map map.json [--size 8.8]

map.json 形如 {"3": ["[1] ROSS D S, et al. Thyroid, 2016, 26(10): 1343-1421."], ...}
键是幻灯片序号（1 起）。会把该页引用的文献拼成一行放在页面底部。
"""
import argparse, json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ap = argparse.ArgumentParser()
ap.add_argument("--deck", required=True)
ap.add_argument("--map", required=True)
ap.add_argument("--size", type=float, default=8.8)
ap.add_argument("--color", default="7A8B99")
ap.add_argument("--bottom-margin", type=float, default=1.10, help="距底距离（英寸）")
a = ap.parse_args()

prs = Presentation(a.deck)
refs = json.load(open(a.map, encoding="utf-8"))
H = prs.slide_height / 914400.0
W = prs.slide_width / 914400.0
color = RGBColor.from_string(a.color.upper())

for key, val in refs.items():
    if isinstance(val, dict):
        items = val.get("items", [])
        margin = val.get("bottom_margin", a.bottom_margin)
        size = val.get("size", a.size)
    else:
        items, margin, size = val, a.bottom_margin, a.size
    i = int(key) - 1
    if not (0 <= i < len(prs.slides)):
        continue
    slide = prs.slides[i]
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(H - margin), Inches(W - 1.1), Inches(0.42 if size >= 8.5 else 0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    p = tf.paragraphs[0]
    p.line_spacing = 1.1
    r1 = p.add_run(); r1.text = "参考文献  "
    r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = color
    r2 = p.add_run(); r2.text = "  ".join(items)
    r2.font.size = Pt(size); r2.font.color.rgb = color

prs.save(a.deck)
print("footers added to", len(refs), "slides")
