# -*- coding: utf-8 -*-
"""
python-docx 排版库 —— PBL 文献综述报告

用法：
    from word_helpers import *
    new_document()
    title_block("研究问题标题", "内分泌系统与疾病 · PBL Case 1")
    h2("摘要")
    body("……")
    ...
    save("xxx.docx")

核心约定（用户硬性要求，不要改）：
  · 全文文字纯黑
  · 不写副标题
  · 各级标题中文黑体 / 西文 Times New Roman
  · 正文中文宋体 / 西文 Times New Roman，12pt，1.5 倍行距，首行缩进 2 字符
"""
import re

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 全黑 —— 不要为了"好看"改回彩色
BLACK = RGBColor(0x00, 0x00, 0x00)
BLACK_DOC = None          # Document 实例，由 new_document() 创建

EN_FONT = "Times New Roman"
CN_BODY = "宋体"
CN_HEAD = "黑体"

THEME_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")


# ============================================================
# 底层：字体设置（含主题属性清理 —— 这一步漏了黑体就不生效）
# ============================================================
def _apply_fonts(rpr, cn, en):
    """在 <w:rPr> 上设置字体，并删掉优先级更高的主题属性。"""
    rf = rpr.get_or_add_rFonts()
    for attr in THEME_ATTRS:                     # ← 关键：主题属性优先于显式属性
        if rf.get(qn(attr)) is not None:
            del rf.attrib[qn(attr)]
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:cs"), en)
    rf.set(qn("w:eastAsia"), cn)


def set_font(run, cn=CN_BODY, en=EN_FONT, size=12, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    _apply_fonts(run._element.get_or_add_rPr(), cn, en)
    run.font.color.rgb = color if color is not None else BLACK


# ============================================================
# 文档初始化
# ============================================================
def new_document(margins=(2.54, 2.54, 3.0, 3.0)):
    """A4 纵向；margins = (上, 下, 左, 右)，单位 cm。"""
    global BLACK_DOC
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(margins[0]), Cm(margins[1])
    sec.left_margin, sec.right_margin = Cm(margins[2]), Cm(margins[3])

    st = doc.styles["Normal"]
    st.font.size = Pt(12)
    st.font.color.rgb = BLACK
    _apply_fonts(st.element.get_or_add_rPr(), CN_BODY, EN_FONT)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(0)

    _style_heading(doc, "Heading 1", 16, 18, 10)
    _style_heading(doc, "Heading 2", 13.5, 12, 7)
    _style_heading(doc, "Heading 3", 12, 9, 5)

    BLACK_DOC = doc
    return doc


def _style_heading(doc, name, size, before, after):
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = BLACK
    _apply_fonts(s.element.get_or_add_rPr(), CN_HEAD, EN_FONT)   # ← 清主题属性
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.line_spacing = 1.35
    s.paragraph_format.keep_with_next = True


# ============================================================
# 段落级 API
# ============================================================
def _h(level, text, size):
    """标题：样式 + 段落 run 双层写死黑体，不依赖继承。"""
    p = BLACK_DOC.add_paragraph(text, style="Heading %d" % level)
    for r in p.runs:
        set_font(r, cn=CN_HEAD, size=size, bold=True, color=BLACK)
    return p


def h1(t):
    return _h(1, t, 16)


def h2(t):
    return _h(2, t, 13.5)


def h3(t):
    return _h(3, t, 12)


def title_block(main, course=None, divider=True):
    """主标题（不写副标题）+ 可选的课程标识行 + 分隔线。"""
    p = BLACK_DOC.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(main), cn=CN_HEAD, size=21, bold=True)

    if course:
        p = BLACK_DOC.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(16)
        set_font(p.add_run(course), cn=CN_BODY, size=10.5)

    if divider:
        p = BLACK_DOC.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "000000")
        bdr.append(bottom)
        pPr.append(bdr)


def body(text, indent=True, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6):
    p = BLACK_DOC.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Pt(size * 2) if indent else Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.5
    set_font(p.add_run(text), size=size)
    return p


def bullet(text, size=11.5):
    p = BLACK_DOC.add_paragraph()
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.first_line_indent = Pt(-12)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.4
    set_font(p.add_run("● " + text), size=size)
    return p


def keywords(text, label="关键词："):
    p = BLACK_DOC.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run(label), cn=CN_HEAD, size=11, bold=True)
    set_font(p.add_run(text), size=11)


def caption(text):
    """图表题注：黑体、居中、小一号。"""
    p = BLACK_DOC.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    set_font(p.add_run(text), cn=CN_HEAD, size=10.5)
    return p


def picture(path, width_cm=15.0):
    """居中插图，题注请紧接着用 caption()。"""
    p = BLACK_DOC.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    return p


# ============================================================
# 表格
# ============================================================
def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def make_table(headers, rows, widths=None, font_size=10.5, header_fill="E9EDF1"):
    """表头浅灰底 + 黑字；跨页重复表头；单行不跨页断开。"""
    t = BLACK_DOC.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, htxt in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        set_font(p.add_run(htxt), cn=CN_HEAD, size=font_size, bold=True, color=BLACK)
        _shade(cell, header_fill)

    for r in rows:
        cells = t.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            set_font(p.add_run(str(val)), size=font_size)

    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)

    # 表头跨页重复
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    # 禁止单行跨页断开
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

    BLACK_DOC.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def ref(text):
    """参考文献条目：悬挂缩进。"""
    p = BLACK_DOC.add_paragraph()
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.first_line_indent = Pt(-24)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), size=10.5)
    return p


def load(path):
    """打开已有 docx，便于用 renumber_refs() 等函数修复。"""
    global BLACK_DOC
    BLACK_DOC = Document(path)
    return BLACK_DOC


def save(path):
    BLACK_DOC.save(path)
    import os
    print("saved:", path, os.path.getsize(path) // 1024, "KB")


# ============================================================
# 参考文献：顺序编码制（按正文首次引用顺序编号）
# ============================================================
_CITE_RE = re.compile(r"\[(\d+(?:\s*[,，、\-–—]\s*\d+)*)\]")


def _expand_cite(group):
    """'3-5,8' -> [3, 4, 5, 8]"""
    nums = []
    for part in re.split(r"[,，、]", group):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^(\d+)\s*[\-–—]\s*(\d+)$", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            nums.extend(range(a, b + 1) if a <= b else [a])
        elif part.isdigit():
            nums.append(int(part))
    return nums


def _compress(nums):
    """[1,2,3,5] -> '1-3,5'"""
    nums = sorted(set(nums))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j == i:
            out.append(str(nums[i]))
        elif j == i + 1:
            out.append("%d,%d" % (nums[i], nums[j]))
        else:
            out.append("%d-%d" % (nums[i], nums[j]))
        i = j + 1
    return ",".join(out)


def _para_text(p):
    return "".join(r.text for r in p.runs)


def _set_para_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def _split_refs(paras, heading):
    for i, p in enumerate(paras):
        if p.text.strip() == heading:
            return i, paras[:i], paras[i + 1:]
    return None, None, None


def _parse_entries(tail):
    entries, ref_paras = {}, []
    for p in tail:
        m = re.match(r"^\s*\[(\d+)\]\s*(.*)$", p.text.strip())
        if m:
            entries[int(m.group(1))] = m.group(2).strip()
            ref_paras.append(p)
    return entries, ref_paras


def _cite_order(body):
    order = []
    for p in body:
        for m in _CITE_RE.finditer(_para_text(p)):
            for n in _expand_cite(m.group(1)):
                if n not in order:
                    order.append(n)
    return order


def renumber_refs(heading="参考文献", verbose=True):
    """
    按正文首次引用顺序重排参考文献编号（GB/T 7714 顺序编码制）。

    正文可以先用任意编号写，写完调用一次即可；编号、区间、逗号列表都会重写。
    未被正文引用的条目排在最后。**必须在 save() 之前调用。**
    返回 {旧编号: 新编号}。
    """
    paras = BLACK_DOC.paragraphs
    ri, body, tail = _split_refs(paras, heading)
    if ri is None:
        raise ValueError("找不到参考文献标题：%r" % heading)

    order = _cite_order(body)
    entries, ref_paras = _parse_entries(tail)
    if not entries:
        raise ValueError("参考文献区没有 [1] 形式的条目")

    missing = [n for n in order if n not in entries]
    if missing:
        raise ValueError("正文引用了不存在的文献编号：%s" % sorted(set(missing)))

    unused = [n for n in sorted(entries) if n not in order]
    final = order + unused
    old2new = {old: i + 1 for i, old in enumerate(final)}

    for p in body:
        txt = _para_text(p)
        if not _CITE_RE.search(txt):
            continue
        _set_para_text(p, _CITE_RE.sub(
            lambda m: "[%s]" % _compress([old2new[n] for n in _expand_cite(m.group(1))]), txt))

    for slot, old in enumerate(final):
        _set_para_text(ref_paras[slot], "[%d] %s" % (slot + 1, entries[old]))

    if verbose:
        changed = sum(1 for o, n in old2new.items() if o != n)
        print("参考文献已重排：%d 条，%d 条编号变化" % (len(final), changed))
        if unused:
            print("  ⚠ 未被正文引用、已排在最后：%s" % unused)
    return old2new


def check_refs(heading="参考文献", verbose=True):
    """校验：正文首次引用顺序必须是 1,2,3…，且与文献表一一对应。"""
    paras = BLACK_DOC.paragraphs
    ri, body, tail = _split_refs(paras, heading)
    if ri is None:
        if verbose:
            print("参考文献检查：未找到「%s」，跳过" % heading)
        return True

    order = _cite_order(body)
    entries, _ = _parse_entries(tail)
    problems = []

    if order != sorted(order):
        problems.append("正文首次引用顺序不是递增：%s" % order)
    if order and order[0] != 1:
        problems.append("第一条引用不是 [1]：%s" % order[0])
    if sorted(entries) != list(range(1, len(entries) + 1)):
        problems.append("文献表编号不连续：%s" % sorted(entries))
    cited, listed = set(order), set(entries)
    if cited - listed:
        problems.append("正文引用了文献表没有的编号：%s" % sorted(cited - listed))
    if listed - cited:
        problems.append("文献表有条目未被正文引用：%s" % sorted(listed - cited))

    # GB/T 7714 基本格式
    for n in sorted(entries):
        txt = entries[n]
        if not re.search(r"\[(J/OL|J|M)\]", txt):
            problems.append("[%d] 缺少 [J]/[M] 文献类型标识" % n)
        elif "[M]" not in txt:
            has_pages = re.search(r"\d{4},\s*[^:]+:\s*[A-Za-z]?\d", txt)
            if not has_pages and "DOI" not in txt.upper():
                problems.append("[%d] 缺卷期页码且无 DOI" % n)

    if verbose:
        print("参考文献检查：", "通过" if not problems else problems)
    return not problems


# ============================================================
# 自检
# ============================================================
def audit():
    """断言：全文纯黑 + 标题黑体 + 无主题字体残留。"""
    from docx.oxml.ns import qn as _qn
    bad_color, bad_head, theme_left = [], [], []
    for p in BLACK_DOC.paragraphs:
        for r in p.runs:
            rPr = r._element.rPr
            if rPr is None:
                continue
            c = rPr.find(_qn("w:color"))
            if c is not None and c.get(_qn("w:val")) != "000000":
                bad_color.append(r.text[:20])
            rf = rPr.find(_qn("w:rFonts"))
            if rf is not None:
                if any(rf.get(_qn(a)) is not None for a in THEME_ATTRS):
                    theme_left.append(r.text[:20])
                if p.style.name.startswith("Heading") and rf.get(_qn("w:eastAsia")) != CN_HEAD:
                    bad_head.append((p.style.name, r.text[:20], rf.get(_qn("w:eastAsia"))))
    print("非黑色文字:", bad_color or "无")
    print("非黑体标题:", bad_head or "无")
    print("主题字体残留:", theme_left or "无")
    ok_refs = check_refs(verbose=True)
    return not (bad_color or bad_head or theme_left) and ok_refs


# ============================================================
# 示例
# ============================================================
if __name__ == "__main__":
    new_document()
    title_block("示例：某因素对某疾病的影响及机制研究", "内分泌系统与疾病 · PBL Case 1")
    h2("摘要")
    body("这里写摘要。所有文字为黑色，标题黑体，西文 Times New Roman[1]。")
    keywords("关键词一；关键词二；关键词三")
    h1("1  引言")
    h2("1.1  案例背景")
    body("这里写案例回顾。")
    make_table(["维度", "典型表现", "本例表现"],
               [["维度一", "该维度的典型表现", "本例的具体数值或描述"]],
               widths=[3.0, 5.0, 7.0])
    caption("表 1  示例表格")
    h1("参考文献")
    ref("[1] Author A, Author B. Title of the article. Journal Name. 2024;1(2):3-4.")
    save("示例.docx")
    audit()
