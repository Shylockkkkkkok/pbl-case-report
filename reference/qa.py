# -*- coding: utf-8 -*-
"""
PPT 版面自查 —— 不启动 Office

用法：
    python qa.py 汇报.pptx                 # 只做几何检查
    python qa.py 汇报.pptx preview/        # 几何检查 + 渲染预览 PNG

检查项：
  · 形状越界（装饰性出血图形会报 OUT-OF-BOUNDS，属已知例外）
  · 文字溢出（按字号逐字估算所需行高）
  · 文字框互相重叠

⚠️ 不要用 PowerPoint COM 导出图片：脚本结尾的 $ppt.Quit() 会关掉整个
   PowerPoint 进程，把用户正在编辑的演示文稿一起关掉。用本脚本代替。

预览渲染是近似实现，以下与真实 PowerPoint 有差异，判断时注意：
  · 英文按字符折行（PowerPoint 按单词）→ 英文可能出现断词
  · 不画箭头头部
  · 字体用微软雅黑代替思源黑体
"""
import io
import os
import re
import sys
import unicodedata

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont

EMU = 914400.0
SCALE = 128                       # 预览：像素 / 英寸
W_IN, H_IN = 13.333, 7.5

FONT_BOLD = "C:/Windows/Fonts/msyhbd.ttc"
FONT_REG = "C:/Windows/Fonts/msyh.ttc"


# ============================================================
def em_width(ch, fs):
    if unicodedata.east_asian_width(ch) in ("W", "F"):
        return fs
    return fs * 0.28 if ch == " " else fs * 0.55


def est_lines(text, box_w_pt, fs):
    if not text.strip():
        return 0
    lines = 0
    for para in text.split("\n"):
        if not para:
            lines += 1
            continue
        cur = 0.0
        for ch in para:
            w = em_width(ch, fs)
            if cur + w > box_w_pt:
                lines += 1
                cur = w
            else:
                cur += w
        lines += 1
    return lines


def run_fs(tf, default=12.0):
    sizes = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
    return max(sizes) if sizes else default


def check(prs):
    problems = []
    for si, slide in enumerate(prs.slides, 1):
        texts = []
        for sh in slide.shapes:
            try:
                l, t = sh.left / EMU, sh.top / EMU
                w, h = sh.width / EMU, sh.height / EMU
            except Exception:
                continue
            if l < -0.02 or t < -0.02 or l + w > W_IN + 0.02 or t + h > H_IN + 0.02:
                problems.append((si, "OUT-OF-BOUNDS", "%.2f,%.2f %.2fx%.2f" % (l, t, w, h), ""))
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text
            if not txt.strip():
                continue
            fs = run_fs(sh.text_frame)
            need = est_lines(txt, w * 72, fs)
            need_h = need * fs * 1.34 / 72.0
            if need_h > h + 0.06:
                problems.append((si, "TEXT-OVERFLOW",
                                 'need %.2f" > box %.2f" (fs=%.1f, lines=%d)' % (need_h, h, fs, need),
                                 txt[:40].replace("\n", "/")))
            texts.append((l, t, w, h, txt[:24].replace("\n", "/")))

        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                a, b = texts[i], texts[j]
                ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
                oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
                if ox > 0.06 and oy > 0.06:
                    problems.append((si, "TEXT-OVERLAP", "%.2fx%.2f\"" % (ox, oy),
                                     "%s || %s" % (a[4], b[4])))
    return problems


# ============================================================
def render(prs, outdir):
    os.makedirs(outdir, exist_ok=True)
    WPX, HPX = int(W_IN * SCALE), int(H_IN * SCALE)
    cache = {}

    def font(sz, bold):
        k = (round(sz), bold)
        if k not in cache:
            cache[k] = ImageFont.truetype(FONT_BOLD if bold else FONT_REG,
                                          max(6, int(round(sz * SCALE / 72.0))))
        return cache[k]

    def rgb(c):
        try:
            return "#" + str(c.rgb)
        except Exception:
            return None

    def px(v):
        return v / EMU * SCALE

    for si, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (WPX, HPX), "white")
        d = ImageDraw.Draw(img)
        try:
            f = slide.background.fill
            if f.type is not None and f.type == 1:
                col = rgb(f.fore_color)
                if col:
                    d.rectangle([0, 0, WPX, HPX], fill=col)
        except Exception:
            pass

        for sh in slide.shapes:
            try:
                l, t, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
            except Exception:
                continue
            st = sh.shape_type
            box = [l, t, l + w, t + h]

            if st == MSO_SHAPE_TYPE.PICTURE:
                try:
                    im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
                    im = im.resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS)
                    img.paste(im, (int(l), int(t)), im)
                except Exception:
                    d.rectangle(box, outline="#cccccc")
                continue

            xml = sh._element.xml

            # pptxgenjs 的线条用 prst="line"，python-pptx 映射不到 auto_shape_type
            if st == MSO_SHAPE_TYPE.AUTO_SHAPE and 'prst="line"' in xml:
                try:
                    col = rgb(sh.line.color) or "#888888"
                    lw = max(1, int(px(sh.line.width or 9525)))
                except Exception:
                    col, lw = "#888888", 1
                fH = fV = False
                try:
                    x = sh._element.spPr.xfrm
                    fH = x.get("flipH") == "1"
                    fV = x.get("flipV") == "1"
                except Exception:
                    pass
                x1, y1, x2, y2 = l, t, l + w, t + h
                if fH:
                    x1, x2 = x2, x1
                if fV:
                    y1, y2 = y2, y1
                d.line([x1, y1, x2, y2], fill=col, width=lw)
                continue

            if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill = line = None
                lw = 1
                try:
                    if sh.fill.type == 1:
                        fill = rgb(sh.fill.fore_color)
                except Exception:
                    pass
                # 空 <a:ln/> 表示无边框；只有真带颜色才画
                m = re.search(r"<a:ln[ >].*?</a:ln>|<a:ln/>", xml, re.S)
                if m and "srgbClr" in m.group(0):
                    try:
                        if sh.line.fill.type == 1:
                            line = rgb(sh.line.color)
                            lw = max(1, int(px(sh.line.width or 9525)))
                    except Exception:
                        pass
                try:
                    name = str(sh.auto_shape_type)
                except Exception:
                    name = ""
                if fill or line:
                    if "OVAL" in name:
                        d.ellipse(box, fill=fill, outline=line, width=lw if line else 1)
                    else:
                        rad = int(min(w, h) * 0.16) if "ROUNDED" in name else 0
                        d.rounded_rectangle(box, radius=rad, fill=fill,
                                            outline=line, width=lw if line else 1)
                if not (sh.has_text_frame and sh.text_frame.text.strip()):
                    continue

            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            if not tf.text.strip():
                continue
            sizes = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
            fs = max(sizes) if sizes else 12
            bold = any(bool(r.font.bold) for p in tf.paragraphs for r in p.runs)
            col = None
            for p in tf.paragraphs:
                for r in p.runs:
                    col = rgb(r.font.color) or col
            col = col or "#000000"
            fnt = font(fs, bold)
            anchor = str(tf.vertical_anchor)
            lh = fs * SCALE / 72.0 * 1.30

            lines = []
            for p in tf.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt:
                    lines.append("")
                    continue
                cur = ""
                for ch in txt:
                    if d.textlength(cur + ch, font=fnt) > w - 2:
                        lines.append(cur)
                        cur = ch
                    else:
                        cur += ch
                lines.append(cur)

            total = len(lines) * lh
            if "MIDDLE" in anchor:
                y0 = t + (h - total) / 2
            elif "BOTTOM" in anchor:
                y0 = t + h - total
            else:
                y0 = t
            al = str(tf.paragraphs[0].alignment) if tf.paragraphs else "None"
            for k, ln in enumerate(lines):
                if not ln:
                    continue
                tw = d.textlength(ln, font=fnt)
                if "CENTER" in al:
                    x0 = l + (w - tw) / 2
                elif "RIGHT" in al:
                    x0 = l + w - tw
                else:
                    x0 = l
                d.text((x0, y0 + k * lh), ln, font=fnt, fill=col)

        img.save(os.path.join(outdir, "s%02d.png" % si))
    print("预览已渲染 ->", outdir)


# ============================================================
def check_refs(prs):
    """页脚引用编号必须按幻灯片顺序首次出现为 1,2,3…，且与文献页一一对应。"""
    ref_idx = None
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() == "参考文献":
                ref_idx = i
                break
        if ref_idx is not None:
            break
    if ref_idx is None:
        return []

    order = []
    for i, s in enumerate(prs.slides):
        if i == ref_idx:
            continue
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for m in re.finditer(r"\[(\d+)\]", sh.text_frame.text):
                n = int(m.group(1))
                if n not in order:
                    order.append(n)

    listed = set()
    for sh in prs.slides[ref_idx].shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            m = re.match(r"^\s*\[(\d+)\]", para.text)
            if m:
                listed.add(int(m.group(1)))

    out = []
    if order != sorted(order):
        out.append("页脚首次引用顺序不是递增: %s" % order)
    if order and order[0] != 1:
        out.append("第一条引用不是 [1]: %s" % order[0])
    missing = [n for n in order if n not in listed]
    if missing:
        out.append("正文引用了文献页没有的编号: %s" % missing)
    unused = sorted(n for n in listed if n not in order)
    if unused:
        out.append("文献页有正文未引用的条目（可接受，但需确认）: %s" % unused)
    return out


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    path = sys.argv[1]
    prs = Presentation(path)

    problems = check(prs)
    print("=" * 70)
    print("几何检查：", path, "（%d 页）" % len(prs.slides))
    print("=" * 70)
    if not problems:
        print("未发现问题。")
    else:
        cur = None
        for si, kind, detail, txt in problems:
            if si != cur:
                print("\n--- 第 %d 页 ---" % si)
                cur = si
            print("  [%s] %s\n      %s" % (kind, detail, txt))
    print("\n合计:", len(problems), "（OUT-OF-BOUNDS 若为装饰性出血图形可忽略）")

    ref_problems = check_refs(prs)
    print("\n参考文献顺序检查：", "通过" if not ref_problems else "有问题")
    for msg in ref_problems:
        print("  ⚠", msg)

    if len(sys.argv) > 2:
        render(prs, sys.argv[2])


if __name__ == "__main__":
    main()
